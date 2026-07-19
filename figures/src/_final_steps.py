# -*- coding: utf-8 -*-
# ขั้นสุดท้าย: xlsx ของทุกตาราง + แพตช์ เสนอ.md / บทที่ 1.md + ตรวจไฟล์ภาพครบ
import io, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

MAP = [('2.1', '_table-extra-1'), ('2.2', '_table-2-2'), ('2.3', '_table-2-html'),
       ('2.4', '_table-2-css'), ('2.5', '_table-2-js'), ('2.6', '_table-2-7'),
       ('2.7', '_table-2-8'), ('2.8', '_table-2-9'), ('2.9', '_table-2-6'),
       ('2.10', '_table-extra-2')]

def cells_from_html(path):
    h = io.open(path, encoding='utf-8').read()
    rows = []
    for tr in re.findall(r'<tr>(.*?)</tr>', h, re.S):
        cs = re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>', tr, re.S)
        cs = [re.sub(r'<[^>]+>', '', c).replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&').strip() for c in cs]
        rows.append(cs)
    return rows

thin = Side(style='thin', color='B7BFD6')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
for num, src in MAP:
    rows = cells_from_html('figures/src/' + src + '.html')
    wb = Workbook(); ws = wb.active; ws.title = 'ตาราง ' + num
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = bd
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if r == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill('solid', fgColor='E9EDF7')
    widths = [28, 30, 60][:len(rows[0])] + [40] * 10
    for c in range(1, len(rows[0]) + 1):
        ws.column_dimensions[chr(64 + c)].width = widths[c - 1] if c <= 3 else 40
    wb.save(f'draw/ตารางที่-{num}.xlsx')
print('xlsx saved:', len(MAP))

# ---- เสนอ.md: ใส่ผังงานจริงแทน placeholder ----
p = 'เสนอ.md'
s = io.open(p, encoding='utf-8').read()
s = s.replace('*(แทรก Flowchart ขั้นตอนการพัฒนาโครงงานที่นี่)*',
  '![ผังงานขั้นตอนการพัฒนาโครงงาน](draw/dev_process_flow.png)\n\n**รูปที่ 12.1 ผังงานขั้นตอนการพัฒนาโครงงาน**')
s = s.replace('*(แทรก Flowchart การทำงานของระบบที่นี่)*',
  '![ผังงานการทำงานของระบบ หน้า 1](draw/mission_flow_page1.png)\n\n**รูปที่ 12.2 ผังงานการทำงานของระบบ (หน้า 1)**\n\n![ผังงานการทำงานของระบบ หน้า 2](draw/mission_flow_page2.png)\n\n**รูปที่ 12.3 ผังงานการทำงานของระบบ (หน้า 2)**')
io.open(p, 'w', encoding='utf-8').write(s)
print('เสนอ.md patched')

# ---- บทที่ 1.md: กรอบแนวคิด ----
p = 'บทที่ 1.md'
s = io.open(p, encoding='utf-8').read()
s = s.replace('*(รูปที่ 1.1 กรอบแนวคิดสำหรับการพัฒนาโครงงาน — แทรกภาพที่นี่)*',
  '![กรอบแนวคิดสำหรับการพัฒนาโครงงาน](draw/conceptual_framework.png)\n\n**รูปที่ 1.1 กรอบแนวคิดสำหรับการพัฒนาโครงงาน**')
io.open(p, 'w', encoding='utf-8').write(s)
print('บทที่ 1.md patched')

# ---- ตรวจว่าไฟล์ภาพที่เอกสารอ้างมีจริงทุกไฟล์ ----
missing = []
for doc in ['บทที่2_เอกสารและงานวิจัยที่เกี่ยวข้อง.md', 'เสนอ.md', 'บทที่ 1.md']:
    s = io.open(doc, encoding='utf-8').read()
    for m in re.findall(r'!\[[^\]]*\]\(([^)]+)\)', s):
        if not os.path.exists(m):
            missing.append(doc + ' -> ' + m)
print('missing refs:', len(missing))
for x in missing: print('  ', x)
